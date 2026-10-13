"""Builds the .xlsx files used by test/import.test.js (run: python3 test/files/make_fixtures.py)."""
import io, os, shutil, zipfile
from openpyxl import load_workbook, Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image

HERE = os.path.dirname(__file__)
TPL = os.path.join(HERE, '..', '..', 'public', 'templates', 'highs-product-import-template.xlsx')

def png(color):
    b = io.BytesIO(); Image.new('RGB', (40, 40), color).save(b, 'PNG'); b.seek(0); return b

# 1) filled template: 2 products + example row; picture on row 4 only
wb = load_workbook(TPL)
s = wb['Supplier']
for r, v in zip(range(4, 11), ['Shenzhen Brightway Electronic Co., Ltd.', 'Bao\'an District, Shenzhen', 'Li Wei',
                                'sales@brightway-led.com', '+86 138 0013 8000', 'liwei_bw', '']):
    s.cell(row=r, column=2, value=v or None)
p = wb['Products']
p.append  # rows 4,5
rows = [
    ['', 'Desk lamp 12W touch dimmer', 'BW-DL12', 'Lighting', '9405.21.00', 0.8, 35, 14, 9, 16, 14.5, 56, 30, 40, 45.0, 'CNY', 300, 'CE'],
    ['', 'Clip lamp 5W', 'BW-CL05', 'Lighting', 'abc', 0.3, '', '', '', 50, 9, 50, 40, 30, '22,50', 'cny', 1000, ''],
]
for i, row in enumerate(rows, start=4):
    for j, v in enumerate(row, start=1):
        if j > 1 and v != '': p.cell(row=i, column=j, value=v)
img = XLImage(png((200, 30, 30))); img.width = img.height = 60
p.add_image(img, 'A4')
wb.save(os.path.join(HERE, 'filled-template.xlsx'))

# 2) supplier's own price list (no template), header on row 3
wb = Workbook(); ws = wb.active; ws.title = 'Price list 2026'
ws['A1'] = 'YIWU SUNHOME HOUSEWARE CO.,LTD  Quotation'
ws.append([])
ws.append(['Picture', 'Item No.', 'Description', 'Size', 'Packing', 'G.W.', 'Meas', 'FOB Price', 'MOQ'])
ws.append(['', 'SH-STB-210', 'Stainless tumbler 500ml', '7x7x20cm', 48, 16, '52*38*23', '$1.85', '1,000'])
ws.append(['', 'SH-MUG-330', 'Ceramic mug 330ml', '9x9x10', 36, 15.5, '45x35x30 cm', '$1.20', 500])
pic = XLImage(png((30, 30, 200))); pic.width = pic.height = 30
ws.add_image(pic, 'A5')
wb.save(os.path.join(HERE, 'price-list.xlsx'))

# 3) WPS-style "embedded in cell" picture via =DISPIMG(...)
src = os.path.join(HERE, 'filled-template.xlsx'); dst = os.path.join(HERE, 'wps-cellimage.xlsx')
zin = zipfile.ZipFile(src); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == 'xl/worksheets/sheet3.xml':
        x = data.decode()
        # put a DISPIMG formula into A5 (row 5 = "Clip lamp")
        import re
        x, n = re.subn(r'<c r="A5"(?: [^>]*?)?(?:/>|>.*?</c>)', '<c r="A5" t="str"><f>_xlfn.DISPIMG("ID_ABC123",1)</f><v>=DISPIMG("ID_ABC123",1)</v></c>', x, count=1)
        assert n == 1, 'A5 not found'
        data = x.encode()
    zout.writestr(item, data)
zout.writestr('xl/cellimages.xml', '<?xml version="1.0" encoding="UTF-8"?><etc:cellImages xmlns:etc="http://www.wps.cn/officeDocument/2017/etCustomData" xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><etc:cellImage><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="2" name="ID_ABC123" descr=""/><xdr:cNvPicPr/></xdr:nvPicPr><xdr:blipFill><a:blip r:embed="rId1"/></xdr:blipFill></xdr:pic></etc:cellImage></etc:cellImages>')
zout.writestr('xl/_rels/cellimages.xml.rels', '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/wps1.jpeg"/></Relationships>')
b = io.BytesIO(); Image.new('RGB', (30, 30), (20, 160, 20)).save(b, 'JPEG'); zout.writestr('xl/media/wps1.jpeg', b.getvalue())
zout.close()
print('ok')
